import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

os.makedirs('presentation', exist_ok=True)

prs = Presentation()
# Set 16:9 widescreen dimensions
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6] # Blank slide

# Brand Color Palette
COLOR_NAVY = RGBColor(30, 58, 138)     # #1E3A8A
COLOR_BLUE = RGBColor(59, 130, 246)    # #3B82F6
COLOR_DARK = RGBColor(15, 23, 42)      # #0F172A
COLOR_LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_GRAY = RGBColor(100, 116, 139)   # #64748B
COLOR_RED = RGBColor(239, 68, 68)      # #EF4444
COLOR_GREEN = RGBColor(16, 185, 129)   # #10B981

def add_header(slide, title_text, category_text="PREDICT & RETAIN — HR WORKFORCE ANALYTICS"):
    # Category tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_BLUE
    
    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_NAVY

def add_card(slide, left, top, width, height, bg_color=COLOR_WHITE, border_color=COLOR_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)
    return shape

# ==============================================================================
# SLIDE 1: Title Slide
# ==============================================================================
slide1 = prs.slides.add_slide(blank_slide_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = COLOR_NAVY
bg1.line.fill.background()

tb = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(3.5))
tf = tb.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
p1.text = "PREDICT & RETAIN"
p1.font.size = Pt(44)
p1.font.bold = True
p1.font.color.rgb = COLOR_WHITE
p1.alignment = PP_ALIGN.LEFT

p2 = tf.add_paragraph()
p2.text = "Data-Driven HR Employee Attrition & Workforce Analytics"
p2.font.size = Pt(24)
p2.font.bold = True
p2.font.color.rgb = COLOR_BLUE
p2.alignment = PP_ALIGN.LEFT
p2.space_before = Pt(10)

p3 = tf.add_paragraph()
p3.text = "Comprehensive Exploratory Analysis, Inferential Validation & Interactive Decision Support"
p3.font.size = Pt(14)
p3.font.color.rgb = RGBColor(224, 231, 255)
p3.alignment = PP_ALIGN.LEFT
p3.space_before = Pt(15)

p4 = tf.add_paragraph()
p4.text = "Domain: Foundations of Data Science  |  October 2026"
p4.font.size = Pt(12)
p4.font.color.rgb = RGBColor(147, 197, 253)
p4.alignment = PP_ALIGN.LEFT
p4.space_before = Pt(25)

# ==============================================================================
# SLIDE 2: Problem Statement
# ==============================================================================
slide2 = prs.slides.add_slide(blank_slide_layout)
add_header(slide2, "Problem Statement: The Hidden Cost of Talent Turnover")

# 3 Column Cards
c_w = Inches(3.6)
c_h = Inches(4.8)
top_pos = Inches(1.8)

# Card 1: Organizational Disruption
add_card(slide2, Inches(0.8), top_pos, c_w, c_h, COLOR_WHITE, COLOR_BLUE)
tb = slide2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Operational Disruption"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(12)
bullets1 = [
    "Voluntary departures cause project bottlenecks and loss of institutional capability.",
    "Replacement expenses typically range between 50% and 200% of annual salary per role.",
    "Increased workload on remaining personnel elevates attrition risk across teams."
]
for b in bullets1:
    p = tf.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# Card 2: Strategic HR Questions
add_card(slide2, Inches(4.8), top_pos, c_w, c_h, COLOR_WHITE, COLOR_BLUE)
tb = slide2.shapes.add_textbox(Inches(5.0), top_pos + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Core Business Dilemmas"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(12)
bullets2 = [
    "Which operating units and job functions suffer disproportional turnover rates?",
    "Does compensation explain attrition, or do acute workplace stressors dominate?",
    "What empirical impact does mandatory overtime have on employee retention?",
    "Can early-tenure vulnerability be identified in the first 24 months?"
]
for b in bullets2:
    p = tf.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# Card 3: The Data-Driven Shift
add_card(slide2, Inches(8.8), top_pos, c_w, c_h, COLOR_WHITE, COLOR_BLUE)
tb = slide2.shapes.add_textbox(Inches(9.0), top_pos + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "The Data Science Shift"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(12)
bullets3 = [
    "Transition from reactive exit interviews to evidence-based workforce analytics.",
    "Rigorous statistical validation (Chi-square, Welch's t-test) of operational hypotheses.",
    "Deploying interactive decision intelligence for proactive HR intervention."
]
for b in bullets3:
    p = tf.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# ==============================================================================
# SLIDE 3: Project Objectives
# ==============================================================================
slide3 = prs.slides.add_slide(blank_slide_layout)
add_header(slide3, "Project Objectives & Analytical Framework")

objectives = [
    ("Objective 1: Quantify Attrition Risk", "Measure baseline and segmented turnover rates across Research & Development, Sales, and Human Resources divisions, as well as 9 distinct organizational job roles.", COLOR_NAVY),
    ("Objective 2: Compensation Architecture Analysis", "Evaluate income distributions, median/mean disparities, and statistical significance tests between departing ($4,787) and retained ($6,833) personnel cohorts.", COLOR_BLUE),
    ("Objective 3: Workplace Stress & Commute Evaluation", "Analyze the observational relationships between mandatory overtime status (30.5% vs 10.4%), commute distance (>20 km vs 0-5 km), and employee turnover.", COLOR_NAVY),
    ("Objective 4: Interactive Dashboard & Decision Support", "Deliver a dynamic Streamlit application with live KPI metrics, multi-criteria filtering (Department, Role, Overtime, Travel), and actionable HR guidance.", COLOR_BLUE)
]

for idx, (title, desc, col) in enumerate(objectives):
    y = Inches(1.8 + idx * 1.3)
    add_card(slide3, Inches(0.8), y, Inches(11.7), Inches(1.15), COLOR_WHITE, col)
    tb = slide3.shapes.add_textbox(Inches(1.1), y + Inches(0.12), Inches(11.1), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = col
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(3)

# ==============================================================================
# SLIDE 4: Dataset Overview
# ==============================================================================
slide4 = prs.slides.add_slide(blank_slide_layout)
add_header(slide4, "Dataset Overview: IBM HR Analytics Benchmark")

# Stats ribbon
stats_data = [
    ("1,470", "Total Employee Records", COLOR_NAVY),
    ("35", "Initial Workforce Variables", COLOR_BLUE),
    ("16.12%", "Baseline Attrition Rate", COLOR_RED),
    ("100%", "Data Completeness (0 Nulls)", COLOR_GREEN)
]

for idx, (val, lbl, color) in enumerate(stats_data):
    x = Inches(0.8 + idx * 2.95)
    add_card(slide4, x, Inches(1.7), Inches(2.8), Inches(1.2), COLOR_WHITE, color)
    tb = slide4.shapes.add_textbox(x + Inches(0.1), Inches(1.8), Inches(2.6), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = val
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = lbl
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_GRAY
    p2.alignment = PP_ALIGN.CENTER

# Detail Box
add_card(slide4, Inches(0.8), Inches(3.2), Inches(11.7), Inches(3.6), COLOR_WHITE, COLOR_BLUE)
tb = slide4.shapes.add_textbox(Inches(1.1), Inches(3.4), Inches(11.1), Inches(3.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Dataset Characteristics & Analytical Fields"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(10)

fields = [
    ("Target Variable:", "Attrition ('Yes' = 237, 'No' = 1,233). Baseline turnover is exactly 16.12%."),
    ("Demographics & Experience:", "Age (18-60), Education, EducationField, TotalWorkingYears, YearsAtCompany (0-40)."),
    ("Organizational Structure:", "Department (R&D: 961, Sales: 446, HR: 63), JobRole (9 distinct titles), JobLevel (1-5)."),
    ("Compensation & Performance:", "MonthlyIncome ($1,009 to $19,999), PercentSalaryHike, PerformanceRating, StockOptionLevel."),
    ("Workplace Stress & Surveys:", "OverTime ('Yes': 416, 'No': 1,054), DistanceFromHome (1-29 km), JobSatisfaction (1-4), WorkLifeBalance (1-4).")
]

for title, desc in fields:
    p = tf.add_paragraph()
    run1 = p.add_run()
    run1.text = title + " "
    run1.font.bold = True
    run1.font.size = Pt(12)
    run1.font.color.rgb = COLOR_NAVY
    run2 = p.add_run()
    run2.text = desc
    run2.font.size = Pt(12)
    run2.font.color.rgb = COLOR_DARK
    p.space_after = Pt(5)

# ==============================================================================
# SLIDE 5: Data Cleaning & Feature Engineering
# ==============================================================================
slide5 = prs.slides.add_slide(blank_slide_layout)
add_header(slide5, "Data Cleaning & Feature Engineering Pipeline")

# Card Left: Data Cleaning
add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0), COLOR_WHITE, COLOR_NAVY)
tb = slide5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Rigorous Data Cleaning Protocol"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(12)

clean_steps = [
    "Zero Missing Values: Confirmed complete 1,470 x 35 records with no null imputations required.",
    "Zero Duplicate Rows: Confirmed unique employee identifiers.",
    "Constant Column Audit: Identified uninformative administrative constants (EmployeeCount=1, Over18='Y', StandardHours=80).",
    "Categorical Validation: Verified standardized label encoding across departments, roles, and travel tiers.",
    "Outlier Preservation: Preserved legitimate high incomes ($20k) and senior tenures (40 yrs) to reflect genuine workforce hierarchy."
]
for s in clean_steps:
    p = tf.add_paragraph()
    p.text = "✔ " + s
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# Card Right: Feature Engineering
add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), COLOR_WHITE, COLOR_BLUE)
tb = slide5.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Feature Engineering Implementations"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_BLUE
p.space_after = Pt(12)

fe_steps = [
    "Attrition_Num: Created binary target mapping (Yes -> 1, No -> 0) enabling mathematical aggregation and correlation.",
    "TenureRatio: Ratio of company longevity relative to chronological age (YearsAtCompany / Age) with safe denominator safeguards.",
    "TenureGroup: Segmented career longevity into 0-2 yrs (probation), 3-5 yrs, 6-10 yrs, and 11+ yrs (veteran).",
    "DistanceBand: Segmented commute distance into 0-5 km, 6-10 km, 11-20 km, and 21+ km.",
    "Final Dataset Validation: Successfully generated and verified data/final_dataset.csv (1,470 rows x 39 columns)."
]
for s in fe_steps:
    p = tf.add_paragraph()
    p.text = "⚡ " + s
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# ==============================================================================
# SLIDE 6: Exploratory Data Analysis Overview
# ==============================================================================
slide6 = prs.slides.add_slide(blank_slide_layout)
add_header(slide6, "Exploratory Data Analysis: Workforce Summary")

# Embed eda_charts.png
if os.path.exists('outputs/eda_charts.png'):
    slide6.shapes.add_picture('outputs/eda_charts.png', Inches(0.8), Inches(1.8), width=Inches(6.8))

# Right Insight Box
add_card(slide6, Inches(7.9), Inches(1.8), Inches(4.6), Inches(5.0), COLOR_WHITE, COLOR_NAVY)
tb = slide6.shapes.add_textbox(Inches(8.1), Inches(2.0), Inches(4.2), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Key Workforce Observations"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(10)

eda_points = [
    "Overall Workforce Distribution: 83.88% of employees remain retained (1,233 staff) against a 16.12% attrition baseline (237 staff).",
    "Departmental Differences: Sales exhibits an observed attrition rate of 20.63%, Human Resources 19.05%, and Research & Development 13.84%.",
    "Role Variation: Highest observed attrition occurs in Sales Representatives (39.76%), Laboratory Technicians (23.94%), and Human Resources (23.08%).",
    "Lowest Turnover Roles: Research Directors (2.50%) and Managers (4.90%) exhibit the lowest observed attrition rates."
]
for pt in eda_points:
    p = tf.add_paragraph()
    p.text = "• " + pt
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# ==============================================================================
# SLIDE 7: Overtime & Compensation Analysis
# ==============================================================================
slide7 = prs.slides.add_slide(blank_slide_layout)
add_header(slide7, "Critical Drivers: Overtime Burden & Compensation Gap")

# Left: Overtime Chart
if os.path.exists('outputs/overtime_attrition.png'):
    slide7.shapes.add_picture('outputs/overtime_attrition.png', Inches(0.8), Inches(1.8), width=Inches(5.7))

# Right: Income Chart
if os.path.exists('outputs/income_attrition.png'):
    slide7.shapes.add_picture('outputs/income_attrition.png', Inches(6.8), Inches(1.8), width=Inches(5.7))

# Bottom Stat Callout Card
add_card(slide7, Inches(0.8), Inches(5.7), Inches(11.7), Inches(1.4), COLOR_LIGHT_BG, COLOR_BLUE)
tb = slide7.shapes.add_textbox(Inches(1.0), Inches(5.75), Inches(11.3), Inches(1.3))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Statistical Hypothesis Testing Results:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

p2 = tf.add_paragraph()
p2.text = "• Overtime Association: Attrition is 30.53% for overtime staff vs 10.44% without overtime (Chi2 = 87.5643, p = 8.16e-21). Statistically highly significant."
p2.font.size = Pt(11)
p2.font.color.rgb = COLOR_DARK

p3 = tf.add_paragraph()
p3.text = "• Compensation Gap: Retained personnel earn mean $6,832.74 (median $5,204) vs $4,787.09 (median $3,202) for departed staff (gap: $2,045.65, Welch t = -7.4826, p = 4.43e-13)."
p3.font.size = Pt(11)
p3.font.color.rgb = COLOR_DARK

p4 = tf.add_paragraph()
p4.text = "• Note: Project spec reference salaries were ~$11,330 vs ~$9,592. Actual calculated values from dataset are $6,833 vs $4,787 (p < 0.001); direction is identical."
p4.font.size = Pt(9.5)
p4.font.italic = True
p4.font.color.rgb = COLOR_GRAY

# ==============================================================================
# SLIDE 8: Satisfaction, Tenure & Commute Distance
# ==============================================================================
slide8 = prs.slides.add_slide(blank_slide_layout)
add_header(slide8, "Workplace Dynamics: Tenure, Sentiment & Commute")

# 3 Images / Panels
w3 = Inches(3.7)
top3 = Inches(1.8)

if os.path.exists('outputs/tenure_attrition.png'):
    slide8.shapes.add_picture('outputs/tenure_attrition.png', Inches(0.8), top3, width=w3)

if os.path.exists('outputs/satisfaction_attrition.png'):
    slide8.shapes.add_picture('outputs/satisfaction_attrition.png', Inches(4.8), top3, width=w3)

if os.path.exists('outputs/distance_attrition.png'):
    slide8.shapes.add_picture('outputs/distance_attrition.png', Inches(8.8), top3, width=w3)

# Bottom Insights
add_card(slide8, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.8), COLOR_WHITE, COLOR_NAVY)
tb = slide8.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.6))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Key Observational Insights:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

p2 = tf.add_paragraph()
p2.text = "1. Early Tenure Window: Employees in years 0–2 exhibit 29.82% attrition, which drops sharply to 13.82% in years 3–5 and 8.13% beyond 11 years."
p2.font.size = Pt(11)
p2.font.color.rgb = COLOR_DARK

p3 = tf.add_paragraph()
p3.text = "2. Sentiment Compounding: Low job satisfaction produces 22.84% attrition (vs 11.33% at level 4). Combined with poor work-life balance, attrition exceeds 40%."
p3.font.size = Pt(11)
p3.font.color.rgb = COLOR_DARK

p4 = tf.add_paragraph()
p4.text = "3. Commute Distance Friction: Commuters traveling >20 km exhibit 22.06% attrition compared to 13.77% for staff residing within 5 km."
p4.font.size = Pt(11)
p4.font.color.rgb = COLOR_DARK

# ==============================================================================
# SLIDE 9: Interactive Streamlit Dashboard
# ==============================================================================
slide9 = prs.slides.add_slide(blank_slide_layout)
add_header(slide9, "Interactive Streamlit Dashboard Architecture")

# Left: 4 Features
add_card(slide9, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0), COLOR_WHITE, COLOR_BLUE)
tb = slide9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Production Dashboard Capabilities"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(10)

dash_feats = [
    "Dynamic KPI Metrics Ribbon: Total Employees, Departed Count, Live Attrition Rate (color-coded vs 16.1% baseline), and Average Monthly Income.",
    "Multi-Dimensional Sidebar Filters: Dynamic cascading dropdowns for Department, Job Role, Overtime Status, and Business Travel.",
    "Zero-Row Safe Execution: Gracefully handles narrow query combinations without breaking or throwing runtime errors.",
    "Execution Command: streamlit run dashboard/app.py"
]
for f in dash_feats:
    p = tf.add_paragraph()
    p.text = "✔ " + f
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# Right: 5 Sections
add_card(slide9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), COLOR_WHITE, COLOR_NAVY)
tb = slide9.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Structured Exploration Modules"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_BLUE
p.space_after = Pt(10)

modules = [
    ("Section A: Workforce Overview", "Interactive donut composition, department comparisons, and sorted job role turnover charts."),
    ("Section B: Compensation & Tenure", "Income distribution boxplots, density histograms, tenure bands, and TenureRatio boxplots."),
    ("Section C: Workplace & Satisfaction", "Overtime rate comparisons, survey score breakdowns, and satisfaction × work-life interaction heatmaps."),
    ("Section D: Commute Analysis", "Distance distribution histograms and distance band turnover evaluations."),
    ("Section E: Key Insights & Evidence", "Dynamically calculated metric summaries and evidence-based HR recommendations.")
]
for mod, desc in modules:
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = "• " + mod + ": "
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_NAVY
    run2 = p.add_run()
    run2.text = desc
    run2.font.size = Pt(10.5)
    run2.font.color.rgb = COLOR_DARK
    p.space_after = Pt(5)

# ==============================================================================
# SLIDE 10: Key Analytical Findings
# ==============================================================================
slide10 = prs.slides.add_slide(blank_slide_layout)
add_header(slide10, "Summary of Key Analytical Findings")

findings_grid = [
    ("1. Overtime Disparity", "Employees working overtime exhibit a 30.53% attrition rate vs 10.44% for non-overtime peers (+20.09 percentage points, Chi2 = 87.56, p < 0.001).", COLOR_RED),
    ("2. Compensation Gap", "Retained staff average $6,832.74 monthly earnings vs $4,787.09 for departed personnel ($2,045.65 gap, Welch t = -7.48, p < 0.001).", COLOR_NAVY),
    ("3. The 0–2 Year Vulnerability", "Nearly 30% of employees with under 2 years of tenure exit the firm (29.82%), stabilizing to 13.82% in years 3–5 and 8.13% after 11 years.", COLOR_BLUE),
    ("4. Role Disparity", "Observed attrition rates are highest in Sales Representatives (39.76%), Laboratory Technicians (23.94%), and Human Resources (23.08%).", COLOR_RED),
    ("5. Sentiment Interaction", "Employees reporting low job satisfaction AND poor work-life balance experience peak attrition rates exceeding 40%.", COLOR_NAVY),
    ("6. Commute Distance", "Commuting beyond 20 km elevates observed attrition to 22.06% compared to 13.77% for staff living within 5 km.", COLOR_BLUE)
]

for idx, (title, desc, color) in enumerate(findings_grid):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8 + col * 5.95)
    y = Inches(1.8 + row * 1.7)
    
    add_card(slide10, x, y, Inches(5.75), Inches(1.5), COLOR_WHITE, color)
    tb = slide10.shapes.add_textbox(x + Inches(0.15), y + Inches(0.1), Inches(5.45), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = color
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(3)

# ==============================================================================
# SLIDE 11: Strategic HR Recommendations
# ==============================================================================
slide11 = prs.slides.add_slide(blank_slide_layout)
add_header(slide11, "Strategic HR Action Plan: Evidence-Based Retention")

recs = [
    ("1. Overtime Governance & Workload Audits", "Establish workload monitoring protocols for staff with recurring OverTime status. Rebalance department headcounts to mitigate prolonged overtime exposure.", COLOR_NAVY),
    ("2. Competitive Compensation Calibration", "Conduct market median salary benchmarking specifically for roles exhibiting elevated turnover (Laboratory Technicians: $3,237/mo, Sales Representatives: $2,626/mo).", COLOR_BLUE),
    ("3. Structured 0–24 Month Retention Program", "Implement 30-60-90-180 day structured milestone reviews and assign experienced mentors to all new hires to navigate the critical early-tenure window.", COLOR_NAVY),
    ("4. Flexible & Hybrid Commute Accommodations", "Provide 2–3 days weekly remote flexibility or transit subsidies for personnel residing >15–20 km away to alleviate commute friction.", COLOR_BLUE),
    ("5. Transparent Career Pathways", "Establish explicit advancement criteria and quarterly skills training to support retention in roles with higher observed turnover.", COLOR_NAVY)
]

for idx, (title, desc, color) in enumerate(recs):
    y = Inches(1.7 + idx * 1.05)
    add_card(slide11, Inches(0.8), y, Inches(11.7), Inches(0.95), COLOR_WHITE, color)
    tb = slide11.shapes.add_textbox(Inches(1.0), y + Inches(0.08), Inches(11.3), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = color
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(2)

# ==============================================================================
# SLIDE 12: Conclusion & Future Scope
# ==============================================================================
slide12 = prs.slides.add_slide(blank_slide_layout)
add_header(slide12, "Conclusion & Future Scope")

# Left: Conclusion
add_card(slide12, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0), COLOR_WHITE, COLOR_NAVY)
tb = slide12.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Concluding Summary"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY
p.space_after = Pt(10)

concl_pts = [
    "Predict & Retain demonstrates that employee attrition is strongly patterned by measurable workplace conditions.",
    "Overtime status, compensation disparities, and the initial 24-month tenure window represent the highest leverage points for retention intervention.",
    "Proactive workforce analytics replaces subjective guesswork with statistically validated talent strategies.",
    "The complete end-to-end framework—from cleaning and EDA to the interactive dashboard—equips leadership with decision intelligence."
]
for pt in concl_pts:
    p = tf.add_paragraph()
    p.text = "✔ " + pt
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

# Right: Future Scope
add_card(slide12, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), COLOR_WHITE, COLOR_BLUE)
tb = slide12.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Future Research & Expansion"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_BLUE
p.space_after = Pt(10)

future_pts = [
    "Longitudinal Survival Analysis: Ingest time-series employee logs to estimate time-to-attrition survival probabilities using Cox Proportional Hazards.",
    "Predictive Machine Learning: Develop calibrated classifier ensembles (XGBoost, Random Forest) with SHAP explainability for individual flight-risk scores.",
    "NLP Sentiment Analysis: Analyze open-text exit interview feedback and anonymous internal pulse surveys to capture cultural nuances.",
    "Enterprise HRIS Integration: Connect live API endpoints to HR platforms (Workday, SAP SuccessFactors) for real-time automated monitoring."
]
for pt in future_pts:
    p = tf.add_paragraph()
    p.text = "🚀 " + pt
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(8)

out_pptx = os.path.join('presentation', 'Predict_and_Retain.pptx')
prs.save(out_pptx)
print(f"Saved {out_pptx} successfully with {len(prs.slides)} slides!")
